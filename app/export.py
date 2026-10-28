"""DOCX export: capability statement, compliance matrix, or full bid packet."""
import io
import json
import re
from typing import Literal

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from fastapi import APIRouter, HTTPException, Query
from fastapi.responses import Response

from app import db, firms, ingest
from app.agent import artifacts
from app.agent.matrix import PAST_PERFORMANCE_RE, load_matrix

router = APIRouter()

MATRIX_WIDTHS = [0.30, 0.12, 0.08, 0.50]
PP_BULLET_WORDS = 8
# firm_capability of a matched past-performance row is built as: "<quote>" for "<client>". or "<quote>".
_PP_ENTRY_RE = re.compile(r'"([^"]+)"(?: for "([^"]+)")?\.')

HEADER = "RFP Copilot | DJ AI Business Consultant LLC"
FOOTER = "Draft for human review. Verify every fact before submission."
ACCENT = RGBColor(0x1F, 0x4E, 0x79)
GREY = RGBColor(0x59, 0x59, 0x59)
BORDER = "BFBFBF"
DOCX_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _new_doc() -> Document:
    doc = Document()
    base = doc.styles["Normal"]
    base.font.name, base.font.size = "Calibri", Pt(10.5)
    base.font.color.rgb = RGBColor(0, 0, 0)
    for name, size in (("Title", 22), ("Heading 1", 14), ("Heading 2", 12)):
        st = doc.styles[name]
        st.font.name, st.font.size, st.font.bold = "Calibri", Pt(size), True
        st.font.color.rgb = RGBColor(0, 0, 0) if name == "Title" else ACCENT
    sec = doc.sections[0]
    for part, text in ((sec.header, HEADER), (sec.footer, FOOTER)):
        p = part.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        run.font.size, run.font.color.rgb = Pt(8.5), GREY
    return doc


def _table(doc: Document, headers: list[str], rows: list[list[str | list[str]]], widths: list[float] | None = None) -> None:
    """A cell value is text, or [summary, *bullets]. widths are fractions of the text width."""
    t = doc.add_table(rows=1, cols=len(headers))
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        for k, v in (("val", "single"), ("sz", "4"), ("space", "0"), ("color", BORDER)):
            el.set(qn(f"w:{k}"), v)
        borders.append(el)
    t._tbl.tblPr.append(borders)
    for cell, h in zip(t.rows[0].cells, headers):
        run = cell.paragraphs[0].add_run(h)
        run.bold, run.font.color.rgb = True, ACCENT
    for row in rows:
        for cell, val in zip(t.add_row().cells, row):
            head, *bullets = [val] if isinstance(val, str) else val
            cell.text = head
            for b in bullets:
                cell.add_paragraph(b, style="List Bullet")
    for r in t.rows:
        for cell in r.cells:
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9)
    repeat = OxmlElement("w:tblHeader")
    t.rows[0]._tr.get_or_add_trPr().append(repeat)
    if widths:
        sec = doc.sections[0]
        total = sec.page_width - sec.left_margin - sec.right_margin
        t.autofit = False
        for r in t.rows:
            for cell, w in zip(r.cells, widths):
                cell.width = int(total * w)


def _bullets(doc: Document, items: list[str]) -> None:
    for i in items:
        doc.add_paragraph(i, style="List Bullet")


def _capability(doc: Document, cap: dict) -> None:
    doc.add_heading("Capability Statement", 1)
    cd = cap["company_data"]
    doc.add_paragraph().add_run(cap["headline"]).bold = True
    doc.add_paragraph(f"{cd['name']} | {cd['location'] or ''}")
    doc.add_paragraph("NAICS: " + ", ".join(cd["naics_codes"]) + " | Certifications: " + ", ".join(cd["certifications"]))
    for title, key in (
        ("Core Competencies", "core_competencies"),
        ("Differentiators", "differentiators"),
        ("Past Performance", "past_performance_summaries"),
    ):
        doc.add_heading(title, 2)
        _bullets(doc, cap[key])


def _pp_draft(r: dict) -> str | list[str]:
    """Compact form of a past-performance draft: a one-line summary and up to three bullets. Full quotes stay in the JSON."""
    is_pp = r["requirement_type"] == "eligibility" and r["section_ref"] != "M" and PAST_PERFORMANCE_RE.search(r["requirement"])
    entries = _PP_ENTRY_RE.findall(r["firm_capability"]) if is_pp else []
    if not entries:
        return r["draft_response"]
    n = len(entries)
    summary = f"{n} matching past-performance {'entry' if n == 1 else 'entries'} documented."
    if r["evidence_needed"]:
        summary += f" {r['evidence_needed']}."
    bullets = []
    for quote, _client in entries[:3]:
        words = quote.split()
        bullets.append(" ".join(words[:PP_BULLET_WORDS]) + ("..." if len(words) > PP_BULLET_WORDS else ""))
    return [summary, *bullets]


def _matrix(doc: Document, matrix: dict) -> None:
    doc.add_heading("Compliance Matrix", 1)
    _table(
        doc,
        ["Requirement", "Type", "Gap", "Draft response"],
        [[r["requirement"], r["requirement_type"], r["gap"], _pp_draft(r)] for r in matrix["rows"]],
        MATRIX_WIDTHS,
    )


def _decision(solicitation_id: str, firm_id: int) -> dict | None:
    with db.session() as conn:
        r = conn.execute(
            "SELECT verdict, score, reasons FROM decisions WHERE solicitation_id = ? AND firm_id = ?"
            " ORDER BY id DESC LIMIT 1",
            (solicitation_id, firm_id),
        ).fetchone()
    return dict(r) if r else None


def build_docx(solicitation_id: str, firm_id: int, kind: str) -> bytes:
    """Raises LookupError when the firm/solicitation is missing or a required piece has not been generated yet."""
    sol, firm = ingest.get_solicitation(solicitation_id), firms.get_firm(firm_id)
    if not sol or not firm:
        raise LookupError("firm or solicitation not found")
    cap = artifacts.load_artifact(solicitation_id, firm_id, "capability") if kind != "matrix" else None
    matrix = load_matrix(solicitation_id, firm_id) if kind != "capability" else None
    narrative = artifacts.load_artifact(solicitation_id, firm_id, "narrative") if kind == "bidpacket" else None
    for name, val, needed in (
        ("capability statement", cap, kind != "matrix"),
        ("matrix", matrix, kind != "capability"),
        ("narrative", narrative, kind == "bidpacket"),
    ):
        if needed and val is None:
            raise LookupError(f"no {name} for this firm and solicitation; generate it first")
    doc = _new_doc()
    doc.add_heading(f"{firm['name']}: {sol['title'] or solicitation_id}", 0)
    if kind == "bidpacket":
        dec = _decision(solicitation_id, firm_id)
        doc.add_heading("Decision Summary", 1)
        if dec:
            doc.add_paragraph(f"Verdict: {dec['verdict']} (score {dec['score']})")
            _bullets(doc, json.loads(dec["reasons"]))
        else:
            doc.add_paragraph("No Go/No-Go decision recorded.")
        doc.add_heading("Readiness", 1)
        doc.add_paragraph(f"Overall readiness: {matrix['summary']['overall_readiness_pct']}%")
    if matrix:
        _matrix(doc, matrix)
    if narrative:
        doc.add_heading("Firm Narrative", 1)
        doc.add_paragraph(narrative)
    if cap:
        _capability(doc, cap)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


@router.get("/agent/export/{solicitation_id}")
def export(solicitation_id: str, firm_id: int = Query(...), kind: Literal["capability", "matrix", "bidpacket"] = Query(...)):
    try:
        data = build_docx(solicitation_id, firm_id, kind)
    except LookupError as exc:
        raise HTTPException(404 if "not found" in str(exc) and "generate" not in str(exc) else 409, str(exc))
    return Response(
        data,
        media_type=DOCX_TYPE,
        headers={"Content-Disposition": f'attachment; filename="{kind}-{solicitation_id[:8]}.docx"'},
    )
