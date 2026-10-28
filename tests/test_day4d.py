import io

from docx import Document
from docx.shared import Pt

from app import db, export, firms
from app.agent import matrix
from app.seed import seed_firms
from tests.test_agent import _final, _llm, _solicitation
from tests.test_day4b import HELP_DESK_ENTRIES, _match, _scope, _set_pp
from tests.test_matrix import GOOD_ROWS, _matrix_final, _row
from tests.conftest import rows

EVAL_ROW = {
    **GOOD_ROWS[1],
    "requirement": "Technical approach is the most important factor, weighted 40 percent",
    "section_ref": "M",
    "requirement_type": "evaluation",
    "draft_response": "Weighting noted; structure the proposal to address it.",
}
PP_L_ROW = {
    **GOOD_ROWS[2],
    "requirement": "Offerors shall provide three references for contracts of similar size and scope",
    "section_ref": "L",
    "requirement_type": "submission",
    "firm_capability": matrix.NO_EVIDENCE,
    "gap": "none",
}
NAICS_ROW = {**GOOD_ROWS[0], "firm_capability": 'NAICS "541512" is in the firm codes.'}
PP_M_ROW = {
    **EVAL_ROW,
    "requirement": "Past performance and relevant experience will be evaluated",
}


def _problems(*rs):
    seed_firms()
    return matrix.check_grounding([matrix.MatrixRow.model_validate(r) if isinstance(r, dict) else r for r in rs], firms.get_firm(1))


# 1. evaluation rows
def test_evaluation_row_is_valid_and_schema_allows_it():
    assert _problems(EVAL_ROW) == []
    enum = matrix.MATRIX_SCHEMA["properties"]["rows"]["items"]["properties"]["requirement_type"]["enum"]
    assert "evaluation" in enum and "evaluation" in matrix.RUBRIC


def test_evaluation_row_with_other_gap_or_capability_is_rejected():
    assert any("evaluation factor" in p and "gap must be 'none'" in p
               for p in _problems({**EVAL_ROW, "gap": "partial", "evidence_needed": "x"}))
    assert any("exactly" in p for p in _problems({**EVAL_ROW, "firm_capability": 'Holds "HUBZone" status.'}))


def test_evaluation_rows_are_excluded_from_counts_and_readiness():
    rs = [matrix.MatrixRow.model_validate(r) for r in (GOOD_ROWS[0], EVAL_ROW)]
    s = matrix.summarize(rs)
    assert s.counts == {"none": 1, "partial": 0, "full": 0} and s.overall_readiness_pct == 100
    assert "evaluation" not in s.counts_by_type
    assert matrix.summarize([rs[1]]).overall_readiness_pct == 0


def test_evaluation_row_persists_through_the_database():
    seed_firms()
    sol_id = _solicitation()
    llm, _ = _llm(_matrix_final("r1", [GOOD_ROWS[0], EVAL_ROW]))
    out = matrix.run_matrix(1, sol_id, llm)
    assert out["matrix"]["rows"][1]["requirement_type"] == "evaluation"
    assert [r["requirement_type"] for r in rows("SELECT requirement_type FROM matrix_rows ORDER BY position")] == ["capability", "evaluation"]
    assert matrix.load_matrix(sol_id, 1) == out["matrix"]


def test_migration_6_keeps_existing_rows(tmp_path):
    conn = db.connect(tmp_path / "old.sqlite3")
    conn.execute("CREATE TABLE schema_migrations (version INTEGER PRIMARY KEY)")
    for version, sql in db.MIGRATIONS[:5]:
        conn.executescript(sql)
        conn.execute("INSERT INTO schema_migrations (version) VALUES (?)", (version,))
    conn.execute(
        "INSERT INTO matrix_rows (created_at, solicitation_id, firm_id, position, requirement, section_ref, firm_capability,"
        " gap, evidence_needed, draft_response, requirement_type) VALUES ('t', 's', 1, 0, 'r', 'C', 'f', 'none', '', 'd', 'submission')"
    )
    db.migrate(conn)
    assert conn.execute("SELECT requirement_type FROM matrix_rows").fetchall()[0][0] == "submission"
    conn.execute(
        "INSERT INTO matrix_rows (created_at, solicitation_id, firm_id, position, requirement, section_ref, firm_capability,"
        " gap, evidence_needed, draft_response, requirement_type) VALUES ('t', 's', 1, 1, 'r', 'M', 'f', 'none', '', 'd', 'evaluation')"
    )
    conn.close()


# 2. past-performance matching only for the references/experience row
def test_section_m_rows_never_trigger_past_performance_matching():
    seed_firms()
    llm, sdk = _llm(_matrix_final("r1", [GOOD_ROWS[0], PP_M_ROW, {**GOOD_ROWS[2]}]))
    out = matrix.run_matrix(1, _solicitation(), llm)
    assert len(sdk.calls) == 1  # no scope extraction, no per-entry matching
    types = [(r["section_ref"], r["requirement_type"], r["gap"]) for r in out["matrix"]["rows"]]
    assert types == [("C", "capability", "none"), ("M", "evaluation", "none"), ("M", "capability", "full")]


def test_references_row_in_section_l_is_still_matched():
    seed_firms()
    _set_pp(HELP_DESK_ENTRIES)
    llm, sdk = _llm(_matrix_final("r1", [PP_M_ROW, PP_L_ROW]), _scope(), *[_match(quote=e["scope"]) for e in HELP_DESK_ENTRIES])
    out = matrix.run_matrix(1, _solicitation(), llm)
    m_row, l_row = out["matrix"]["rows"]
    assert (m_row["requirement_type"], m_row["firm_capability"]) == ("evaluation", matrix.NO_EVIDENCE)
    assert (l_row["requirement_type"], l_row["gap"]) == ("eligibility", "none")
    assert len(sdk.calls) == 5


def test_bare_references_wording_is_detected():
    # regression: a stray control character once stopped "references" alone from matching
    assert matrix.PAST_PERFORMANCE_RE.search("Provide three references")


# 3. compact docx
def _pp_matrix_doc(n_matches: int):
    seed_firms()
    _set_pp(HELP_DESK_ENTRIES)
    sol_id = _solicitation()
    ms = [_match(quote=e["scope"]) if i < n_matches else _match(False, (), "") for i, e in enumerate(HELP_DESK_ENTRIES)]
    llm, _ = _llm(_matrix_final("r1", [NAICS_ROW, EVAL_ROW, PP_L_ROW]), _scope(), *ms)
    out = matrix.run_matrix(1, sol_id, llm)
    capability = {"headline": "h", "core_competencies": [], "differentiators": [], "past_performance_summaries": [],
                  "company_data": {"name": "n", "location": "l", "naics_codes": [], "certifications": []}}
    from app.agent import artifacts
    artifacts.save_artifact(sol_id, 1, "capability", capability)
    return out, Document(io.BytesIO(export.build_docx(sol_id, 1, "matrix")))


def test_docx_matrix_table_widths_font_and_repeating_header():
    _, doc = _pp_matrix_doc(3)
    table = doc.tables[0]
    total = doc.sections[0].page_width - doc.sections[0].left_margin - doc.sections[0].right_margin
    widths = [c.width / total for c in table.rows[1].cells]
    assert [round(w, 2) for w in widths] == [0.30, 0.12, 0.08, 0.50]
    assert table.rows[0]._tr.trPr.find(export.qn("w:tblHeader")) is not None
    sizes = {run.font.size for row in table.rows for cell in row.cells for p in cell.paragraphs for run in p.runs}
    assert sizes == {Pt(9)}


def test_docx_past_performance_row_is_compact_and_json_keeps_full_quotes():
    out, doc = _pp_matrix_doc(3)
    cell = doc.tables[0].rows[3].cells[3]
    assert [p.text for p in cell.paragraphs] == [
        "3 matching past-performance entries documented.",
        "Help desk operation and ITSM migration for county...",
        "Help desk operation and ITSM migration for county...",
        "Help desk operation and ITSM migration for county...",
    ]
    assert all(p.style.name == "List Bullet" for p in cell.paragraphs[1:])
    assert all(f'"{e["scope"]}"' in out["matrix"]["rows"][2]["firm_capability"] for e in HELP_DESK_ENTRIES)
    # non past-performance rows keep their full draft text
    assert doc.tables[0].rows[2].cells[3].text == EVAL_ROW["draft_response"]


def test_docx_partial_match_summary_includes_what_is_still_needed():
    _, doc = _pp_matrix_doc(1)
    paras = [p.text for p in doc.tables[0].rows[3].cells[3].paragraphs]
    assert paras[0].startswith("1 matching past-performance entry documented. Provide 2 more")
    assert len(paras) == 2


# 4. "Based in", never "office"
def test_office_wording_is_rejected_unless_the_profile_uses_it():
    assert any("'office'" in p for p in _problems({**GOOD_ROWS[0], "draft_response": "DJ AI maintains an office in Syracuse."}))
    assert _problems({**GOOD_ROWS[0], "draft_response": "Based in Syracuse, NY."}) == []
    seed_firms()
    firm = {**firms.get_firm(1), "name": "Office Works LLC"}
    row = _row(draft_response="DJ AI maintains an office in Syracuse.")
    assert matrix.check_grounding([row], firm) == []


def test_prompt_requires_based_in_wording():
    assert 'write "Based in" followed by the profile\'s place_of_performance' in matrix.RUBRIC
