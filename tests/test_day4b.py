import io
import json

import pytest
from docx import Document
from fastapi.testclient import TestClient

from app import db, firms
from app.agent import artifacts, capability, matrix, narrative, tools
from app.main import app
from app.routes import get_llm
from app.seed import seed_firms
from tests.test_agent import _final, _llm, _message, _response, _solicitation
from tests.test_matrix import GOOD_ROWS, _matrix_final, _pp_none

PP_ROW = {
    **GOOD_ROWS[2],
    "requirement": "Offerors shall provide three references for contracts of similar size and scope",
    "requirement_type": "submission",
    "firm_capability": matrix.NO_EVIDENCE,
    "gap": "none",
}


HELP_DESK_SOL = {"section_c": "The contractor shall provide help desk support and ITSM migration.", "title": "t"}
HELP_DESK_ENTRIES = [
    {"client": f"County {i}", "scope": f"Help desk operation and ITSM migration for county agency {i}", "dates": "2024"}
    for i in range(3)
]


def _set_pp(entries: list[dict]) -> None:
    with db.session() as conn:
        conn.execute("UPDATE firms SET past_performance = ? WHERE id = 1", (json.dumps(entries),))


def _scope(size: str = ""):
    return _final("s", {"terms": ["help desk", "ITSM migration", "SLA reporting"], "size_indicator": size})


def _match(relevant=True, terms=("help desk", "ITSM migration"), quote=""):
    return _final("m", {"relevant": relevant, "matched_terms": list(terms), "quote": quote})


def _assess(llm, sol=HELP_DESK_SOL):
    row = matrix.MatrixRow.model_validate(PP_ROW)
    pp = matrix.assess_past_performance(firms.get_firm(1), sol, llm)
    [out] = matrix.apply_past_performance([row], pp)
    assert matrix.check_grounding([out], firms.get_firm(1)) == []
    return out


def test_seed_profile_against_help_desk_solicitation_is_full():
    seed_firms()
    llm, sdk = _llm(_scope(), *[_match(False, (), "") for _ in range(3)])
    out = _assess(llm)
    assert (out.requirement_type, out.gap, out.firm_capability) == ("eligibility", "full", matrix.NO_EVIDENCE)
    assert out.draft_response.startswith("Gap:")
    assert all(c["reasoning"] == {"effort": "minimal"} for c in sdk.calls) and len(sdk.calls) == 4


def test_fabricated_quote_is_rejected():
    seed_firms()
    fake = _match(True, ("help desk", "ITSM migration"), "Help desk operation and ITSM migration")
    out = _assess(_llm(_scope(), fake, fake, fake)[0])
    assert (out.gap, out.firm_capability) == ("full", matrix.NO_EVIDENCE)


def test_single_term_match_is_rejected():
    seed_firms()
    _set_pp(HELP_DESK_ENTRIES)
    one = [_match(True, ("help desk",), e["scope"]) for e in HELP_DESK_ENTRIES]
    assert _assess(_llm(_scope(), *one)[0]).gap == "full"


def test_three_help_desk_entries_yield_none():
    seed_firms()
    _set_pp(HELP_DESK_ENTRIES)
    out = _assess(_llm(_scope(), *[_match(quote=e["scope"]) for e in HELP_DESK_ENTRIES])[0])
    assert out.gap == "none" and out.evidence_needed == ""
    assert all(f'"{e["scope"]}"' in out.firm_capability for e in HELP_DESK_ENTRIES)


@pytest.mark.parametrize("n,gap", [(1, "partial"), (2, "partial")])
def test_one_or_two_verified_matches_are_partial(n, gap):
    seed_firms()
    _set_pp(HELP_DESK_ENTRIES)
    ms = [_match(quote=e["scope"]) if i < n else _match(False, (), "") for i, e in enumerate(HELP_DESK_ENTRIES)]
    assert _assess(_llm(_scope(), *ms)[0]).gap == gap


def test_size_indicator_downgrades_none_unless_an_entry_is_comparable():
    seed_firms()
    _set_pp(HELP_DESK_ENTRIES)
    ms = [_match(quote=e["scope"]) for e in HELP_DESK_ENTRIES]
    out = _assess(_llm(_scope("1,200 users"), *ms)[0])
    assert out.gap == "partial" and "comparable size (1,200 users)" in out.evidence_needed
    big = [{**e, "scope": e["scope"] + " serving 1,500 users"} for e in HELP_DESK_ENTRIES]
    _set_pp(big)
    ms = [_match(quote=e["scope"]) for e in big]
    assert _assess(_llm(_scope("1,200 users"), *ms)[0]).gap == "none"


def test_scope_extraction_drops_generic_terms():
    terms = {"terms": ["services", "County support", "help desk", "Help Desk", "system"], "size_indicator": " "}
    assert matrix.extract_scope(HELP_DESK_SOL, _llm(_final("s", terms))[0]) == (["help desk"], "")


def test_stopwords_never_reach_the_search():
    seed_firms()
    assert tools.search_past_performance(1, ["services", "System", "county", "management"])["matches"] == []
    [hit] = tools.search_past_performance(1, ["system", "grant"])["matches"]
    assert hit["matched_keywords"] == ["grant"]


def test_run_matrix_overrides_model_gap_for_past_performance_row():
    seed_firms()
    llm, _ = _llm(_matrix_final("r1", [PP_ROW]), _scope(), *[_match(False, (), "") for _ in range(3)])
    out = matrix.run_matrix(1, _solicitation(), llm)
    row = out["matrix"]["rows"][0]
    assert (row["requirement_type"], row["gap"]) == ("eligibility", "full")
    assert out["matrix"]["summary"]["counts_by_type"] == {"capability": 0, "eligibility": 1}


def test_validator_rejects_submission_row_mentioning_past_performance():
    seed_firms()
    problems = matrix.check_grounding([matrix.MatrixRow.model_validate(PP_ROW)], firms.get_firm(1))
    assert any("must be 'eligibility'" in p for p in problems)


def _setup_matrix():
    seed_firms()
    sol_id = _solicitation()
    llm, _ = _llm(_matrix_final("m", GOOD_ROWS), *_pp_none())
    matrix.run_matrix(1, sol_id, llm)
    return sol_id


def _text(words: int, extra: str = "") -> str:
    return ("Syracuse consulting " * words).strip() + extra


def test_narrative_saves_and_prompt_lists_full_gaps():
    sol_id = _setup_matrix()
    text = _text(100)
    llm, sdk = _llm(_response("n", _message(text)))
    assert narrative.run_narrative(1, sol_id, llm) == {"narrative": text}
    call = sdk.calls[0]
    assert call["reasoning"] == {"effort": "minimal"}
    assert "Past performance evaluated on federal DoD contracts" in call["input"][0]["content"]
    assert artifacts.load_artifact(sol_id, 1, "narrative") == text


def test_narrative_number_guard_rejects_invented_numbers_then_retries():
    sol_id = _setup_matrix()
    bad = _text(100, " We delivered 47 projects worth $2M.")
    llm, sdk = _llm(_response("n1", _message(bad)), _response("n2", _message(_text(100, " Since 2023."))))
    out = narrative.run_narrative(1, sol_id, llm)
    assert len(sdk.calls) == 2 and "numbers not in the firm profile" in sdk.calls[1]["input"][0]["content"]
    assert "47" not in out["narrative"]


def test_narrative_requires_matrix_and_fails_after_second_bad_answer():
    seed_firms()
    with pytest.raises(LookupError):
        narrative.run_narrative(1, _solicitation(), _llm()[0])
    sol_id = _setup_matrix()
    bad = _text(100, " 47")
    with pytest.raises(RuntimeError, match="did not pass validation"):
        narrative.run_narrative(1, sol_id, _llm(_response("a", _message(bad)), _response("b", _message(bad)))[0])


def test_number_guard_accepts_profile_numbers():
    seed_firms()
    firm = firms.get_firm(1)
    assert artifacts.ungrounded_numbers("NAICS 541512, $10K-$75K, since 2023.", firm) == []
    assert artifacts.ungrounded_numbers("Won $5,000,000.", firm) == ["$5,000,000"]


STATEMENT = {
    "headline": "AI advisory for CNY organizations",
    "core_competencies": ["Workflow automation"],
    "differentiators": ["HUBZone-certified"],
    "past_performance_summaries": ["IT consulting since 2023."],
}


def test_capability_statement_company_data_filled_by_code():
    seed_firms()
    llm, sdk = _llm(_final("c", STATEMENT))
    out = capability.run_capability(1, _solicitation(), llm)
    assert "company_data" not in sdk.calls[0]["text"]["format"]["schema"]["properties"]
    assert out["company_data"] == {
        "name": "DJ AI Business Consultant LLC",
        "naics_codes": ["541511", "541512", "541611", "518210"],
        "certifications": ["HUBZone"],
        "location": "Syracuse, NY",
    }
    assert out["headline"] == STATEMENT["headline"]


def test_capability_statement_rejects_invented_numbers():
    seed_firms()
    bad = {**STATEMENT, "differentiators": ["Served 300 clients"]}
    llm, sdk = _llm(_final("c1", bad), _final("c2", STATEMENT))
    capability.run_capability(1, _solicitation(), llm)
    assert len(sdk.calls) == 2 and "300" in sdk.calls[1]["input"][0]["content"]


def test_docx_exports_open_and_contain_matrix_table():
    sol_id = _setup_matrix()
    llm, _ = _llm(_response("n", _message(_text(100))), _final("c", STATEMENT))
    app.dependency_overrides[get_llm] = lambda: llm
    try:
        client = TestClient(app)
        params = {"firm_id": 1}
        assert client.get(f"/agent/export/{sol_id}", params={**params, "kind": "bidpacket"}).status_code == 409
        assert client.post("/agent/narrative", json={"firm_id": 1, "solicitation_id": sol_id}).status_code == 200
        assert client.post("/agent/capability-statement", json={"firm_id": 1, "solicitation_id": sol_id}).status_code == 200
        for kind in ("capability", "matrix", "bidpacket"):
            r = client.get(f"/agent/export/{sol_id}", params={**params, "kind": kind})
            assert r.status_code == 200 and "attachment" in r.headers["content-disposition"]
            doc = Document(io.BytesIO(r.content))
            assert doc.sections[0].header.paragraphs[0].text == "RFP Copilot | DJ AI Business Consultant LLC"
            assert doc.sections[0].footer.paragraphs[0].text == "Draft for human review. Verify every fact before submission."
            if kind != "capability":
                table = doc.tables[0]
                assert [c.text for c in table.rows[0].cells] == ["Requirement", "Type", "Gap", "Draft response"]
                assert len(table.rows) == 1 + len(GOOD_ROWS)
            if kind == "bidpacket":
                text = "\n".join(p.text for p in doc.paragraphs)
                assert "Decision Summary" in text and "Overall readiness: 50%" in text and "Firm Narrative" in text
        assert client.get(f"/agent/export/{sol_id}", params={**params, "kind": "bogus"}).status_code == 422
        assert client.get("/agent/export/nope", params={**params, "kind": "matrix"}).status_code == 404
    finally:
        app.dependency_overrides.clear()
